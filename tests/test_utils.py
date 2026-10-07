"""Unit tests for utility functions"""

import os
import tempfile
from app.models import Student
from app.utils import (
    save_students_to_file,
    load_students_from_file,
    format_student_name,
    validate_student_id
)

def test_format_student_name():
    assert format_student_name("ahmed al-kathiri") == "Ahmed Al-kathiri"
    assert format_student_name("  FATIMA   SHANFARI  ") == "Fatima Shanfari"
    assert format_student_name("") == ""

def test_validate_student_id():
    assert validate_student_id("20211001") is True
    assert validate_student_id("12345") is True
    assert validate_student_id("123") is False
    assert validate_student_id("") is False

def test_save_and_load_students(tmp_path):
    temp_file = str(tmp_path / "test_students.json")
    students = [
        Student(id="20211001", name="Ahmed Al-Kathiri", gpa=3.65),
        Student(id="20221045", name="Fatima Al-Shanfari", gpa=3.82)
    ]
    
    save_students_to_file(students, filepath=temp_file)
    assert os.path.exists(temp_file)
    
    loaded = load_students_from_file(filepath=temp_file)
    assert len(loaded) == 2
    assert loaded[0].id == "20211001"
    assert loaded[0].name == "Ahmed Al-Kathiri"
    assert loaded[1].gpa == 3.82

def test_generate_students_at_risk_report_excel(tmp_path):
    from app.utils import generate_students_at_risk_report_excel, get_students_at_risk_breakdown
    import openpyxl
    
    excel_file = str(tmp_path / "test_report.xlsx")
    breakdown = get_students_at_risk_breakdown()
    out = generate_students_at_risk_report_excel(excel_file, breakdown=breakdown)
    
    assert os.path.exists(excel_file)
    assert os.path.getsize(excel_file) > 1000
    
    wb = openpyxl.load_workbook(excel_file)
    assert "Students at Risk Report" in wb.sheetnames
    assert "All At-Risk Data (Filterable)" in wb.sheetnames
    ws1 = wb["Students at Risk Report"]
    assert "DHOFAR UNIVERSITY" in str(ws1['A1'].value)
    assert "STUDENTS AT ACADEMIC RISK" in str(ws1['A2'].value)

def test_load_and_search_policies(tmp_path):
    from app.utils import load_policies_from_file, search_policy_documents, delete_policy_file
    
    p_dir = str(tmp_path / "policies")
    p_json = str(tmp_path / "policies.json")
    
    policies = load_policies_from_file(filepath=p_json, policies_dir=p_dir)
    assert len(policies) >= 2
    assert any("Probation" in p.title for p in policies)
    
    # Search
    results = search_policy_documents("probation 12 credit", policies_dir=p_dir, json_path=p_json)
    assert len(results) > 0
    assert any(len(r['matches']) > 0 for r in results)
    
    # Delete
    del_id = policies[0].id
    assert delete_policy_file(del_id, policies_dir=p_dir, json_path=p_json) is True
    reloaded = load_policies_from_file(filepath=p_json, policies_dir=p_dir)
    assert len(reloaded) == len(policies) - 1


